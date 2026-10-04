"""Regenerates the demo files in samples/. Needs: pip install reportlab pillow python-docx"""
import random, shutil, io
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from PIL import Image as PILImage, ImageDraw, ImageFont
import docx

st = getSampleStyleSheet()
random.seed(7)

# 1) Two short versions (copied from the test fixtures so tests and demo agree)
shutil.copy("tests/fixtures/contract_v1.pdf", "samples/services_agreement_v1.pdf")
shutil.copy("tests/fixtures/contract_v2.pdf", "samples/services_agreement_v2.pdf")

# 2) The same agreement as a Word file (for the DOCX path and highlighting in HTML)
d = docx.Document()
d.add_heading("SERVICES AGREEMENT", 0)
d.add_paragraph("This Agreement is made between Alpha Trading LLC (the Supplier) and Beta Holdings Ltd (the Customer).")
for h, t in [
    ("1. Definitions", "In this Agreement, Services means the services described in Schedule 1 and Fees means the amounts payable under clause 2."),
    ("2. Payment", "The Customer shall pay each invoice within 30 days of receipt. Late payments accrue interest at 2% per month."),
    ("3. Liability", "The Supplier's total liability under this Agreement shall not exceed AED 100,000 in any calendar year."),
    ("4. Termination", "Either party may terminate this Agreement by giving 60 days written notice to the other party."),
    ("5. Confidentiality", "Each party shall keep the other party's Confidential Information secret for a period of five years after termination."),
    ("6. Publicity", "Neither party shall issue any press release about this Agreement without the other party's prior written consent."),
    ("7. Governing Law", "This Agreement is governed by the laws of the United Arab Emirates."),
]:
    d.add_paragraph(h, style="Heading 2"); d.add_paragraph(t)
d.save("samples/services_agreement_v1.docx")

# 3) A scanned PDF: pages are pictures, there is no text layer
img = PILImage.new("RGB", (1240, 1754), "white")
dr = ImageDraw.Draw(img)
try: font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34)
except Exception: font = ImageFont.load_default()
y = 120
for line in ["SERVICES AGREEMENT", "", "1. Payment", "The Customer shall pay each invoice within 30 days.", "", "2. Liability", "Total liability shall not exceed AED 100,000."]:
    dr.text((110, y), line, fill="black", font=font); y += 60
buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
c = canvas.Canvas("samples/scanned_no_text.pdf", pagesize=A4)
c.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1]); c.save()

# 4) A long master agreement (~150 pages) with needles buried deep, for the large-document path
TOPICS = ["Delivery", "Acceptance", "Audit Rights", "Subcontracting", "Insurance", "Force Majeure", "Notices", "Records", "Training",
          "Service Levels", "Change Control", "Escalation", "Warranties", "Intellectual Property", "Publicity", "Assignment",
          "Waiver", "Severability", "Counterparts", "Business Continuity"]
FILL = ["The parties shall cooperate in good faith to give effect to the provisions of this clause and shall each bear their own costs in doing so.",
        "Any notice given under this clause must be in writing and delivered to the address stated in the relevant Statement of Work.",
        "Neither party shall be deemed in breach of this clause to the extent that performance is prevented by circumstances beyond its reasonable control.",
        "The Supplier shall maintain complete and accurate records relating to the Services and make them available on reasonable request.",
        "Where a Service Level is missed the Supplier shall promptly investigate, report on the cause and propose a remediation plan to the Customer.",
        "References to a statute include that statute as amended or re-enacted from time to time and any subordinate legislation made under it."]
NEEDLES = {
    340: ("Data Breach Notification", "The Supplier shall notify the Customer of any Personal Data breach without undue delay and in any event within 72 hours of becoming aware of it."),
    450: ("Limitation of Liability", "The Supplier's aggregate liability arising under this Agreement shall not exceed AED 500,000 in respect of all claims in any contract year."),
    590: ("Early Termination Fee", "If the Customer terminates this Agreement for convenience before the end of the Initial Term it shall pay an early termination fee equal to 25% of the remaining Fees."),
}
doc = SimpleDocTemplate("samples/master_services_agreement_long.pdf", pagesize=A4, leftMargin=70, rightMargin=70, topMargin=70, bottomMargin=70)
body = [Paragraph("MASTER SERVICES AGREEMENT", st["Title"]),
        Paragraph("This Master Services Agreement is made between Alpha Trading LLC (the Supplier) and Beta Holdings Ltd (the Customer).", st["BodyText"]), Spacer(1, 10)]
for n in range(1, 761):
    h, t = NEEDLES.get(n, (f"{random.choice(TOPICS)} ({n})", None))
    body.append(Paragraph(f"{n}. {h}", st["Heading3"]))
    if t: body.append(Paragraph(t, st["BodyText"]))
    for _ in range(random.randint(3, 4)): body.append(Paragraph(random.choice(FILL), st["BodyText"]))
doc.build(body)

open("samples/not_a_contract.txt", "w").write("This plain text file should be rejected: only PDF and DOCX are accepted.\n")
print("done")
